from pathlib import Path
import csv, json, re, sys
from docx import Document

root = Path(__file__).resolve().parents[1]
model = json.loads((root / 'docs/modelo_canonico.json').read_text(encoding='utf-8'))
schema = (root / 'database/01_schema.sql').read_text(encoding='utf-8')
dictionary = (root / 'docs/03_dicionario_dados.md').read_text(encoding='utf-8')
der = (root / 'docs/der/der_radar_urbano.mmd').read_text(encoding='utf-8')
matrix = (root / 'docs/07_matriz_banco_csv.md').read_text(encoding='utf-8')
report_path = root / 'Relatorio_Final_Radar_Urbano_Definitivo.docx'
csv_path = root / 'data/radar_urbano_ocorrencias_teste.csv'
errors = []

report_text = ''
if report_path.exists():
    report = Document(report_path)
    report_parts = [p.text for p in report.paragraphs]
    for report_table in report.tables:
        for row in report_table.rows:
            report_parts.extend(cell.text for cell in row.cells)
    report_text = '\n'.join(report_parts)
else:
    errors.append('Relatório Word definitivo ausente')

def table_block(name):
    m = re.search(rf'CREATE TABLE\s+{re.escape(name)}\s*\((.*?)\n\);', schema, re.I | re.S)
    return m.group(1) if m else None

for table in model['tables']:
    block = table_block(table['name'])
    if block is None:
        errors.append(f"Tabela ausente no DDL: {table['name']}")
        continue
    for field in table['fields']:
        if not re.search(rf'^\s*{re.escape(field["name"])}\s+', block, re.M):
            errors.append(f"Campo ausente no DDL: {table['name']}.{field['name']}")
        for label, content in [('dicionário', dictionary), ('DER', der), ('relatório', report_text)]:
            if field['name'] not in content:
                errors.append(f"Campo ausente no {label}: {table['name']}.{field['name']}")
    if table['name'] not in dictionary or table['name'].upper() not in der or table['name'] not in report_text:
        errors.append(f"Tabela não rastreada em todos os documentos: {table['name']}")

rows = list(csv.DictReader(csv_path.open(encoding='utf-8')))
expected_columns = [c['name'] for c in model['csv_columns']]
if list(rows[0].keys()) != expected_columns:
    errors.append('Cabeçalho do CSV diverge do modelo canônico')
for column in expected_columns:
    if column not in matrix:
        errors.append(f'Coluna ausente na matriz Banco × CSV: {column}')
view_match = re.search(r'CREATE VIEW\s+vw_ocorrencia_exportacao_ia\s+AS\s+SELECT(.*?)\nFROM\s+ocorrencia', schema, re.I | re.S)
if not view_match:
    errors.append('View de exportação ausente no DDL')
else:
    view_columns = []
    for raw in view_match.group(1).split(',\n'):
        line = raw.strip()
        alias = re.search(r'\s+AS\s+(\w+)\s*$', line, re.I)
        direct = re.search(r'\.(\w+)\s*$', line)
        name = alias.group(1) if alias else (direct.group(1) if direct else None)
        if name:
            view_columns.append(name)
    if view_columns != expected_columns:
        errors.append(f'Colunas da view divergem do CSV: {view_columns}')
if len(rows) != 20:
    errors.append(f'Quantidade de registros diferente de 20: {len(rows)}')
ids = [int(r['id_ocorrencia']) for r in rows]
if len(set(ids)) != len(ids): errors.append('IDs duplicados no CSV')
protocols = [r['protocolo'] for r in rows]
if len(set(protocols)) != len(protocols): errors.append('Protocolos duplicados no CSV')
for r in rows:
    if r['prioridade'] not in model['domains']['prioridade']: errors.append(f"Prioridade inválida no ID {r['id_ocorrencia']}")
    if r['prioridade_sugerida'] not in model['domains']['prioridade']: errors.append(f"Prioridade sugerida inválida no ID {r['id_ocorrencia']}")
    if r['categoria_codigo'] not in model['domains']['categoria_codigo']: errors.append(f"Categoria inválida no ID {r['id_ocorrencia']}")
    if r['status_codigo'] not in model['domains']['status_codigo']: errors.append(f"Status inválido no ID {r['id_ocorrencia']}")
    if not -90 <= float(r['latitude']) <= 90 or not -180 <= float(r['longitude']) <= 180: errors.append(f"Coordenadas inválidas no ID {r['id_ocorrencia']}")
    dup = r['id_ocorrencia_duplicada'].strip()
    if dup and (int(dup) not in ids or int(dup) == int(r['id_ocorrencia'])): errors.append(f"Duplicidade inválida no ID {r['id_ocorrencia']}")

high = [int(r['id_ocorrencia']) for r in rows if r['prioridade'] in ('ALTA','CRITICA')]
if high != [3,7,11,14,18]: errors.append(f'Gabarito alta/crítica divergente: {high}')
csv_files = list(root.rglob('*.csv'))
if len(csv_files) != 1: errors.append(f'O pacote deve ter um único CSV; encontrados {len(csv_files)}')

if errors:
    print('VALIDAÇÃO COM ERROS')
    for e in errors: print('- ' + e)
    sys.exit(1)
print('VALIDAÇÃO APROVADA')
print(f'Tabelas verificadas: {len(model["tables"])}')
print(f'Campos verificados: {sum(len(t["fields"]) for t in model["tables"])}')
print('Dicionário, DER e relatório: nomes rastreados')
print('View de exportação: cabeçalho idêntico ao CSV')
print(f'Colunas CSV: {len(expected_columns)}')
print(f'Registros CSV: {len(rows)}')
print('IDs ALTA/CRITICA: 3, 7, 11, 14 e 18')
print('Duplicidade de teste: ID 20 referencia ID 5')
